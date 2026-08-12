#!/usr/bin/env python3
"""
05_assemble.py  —  STAGES D + E: Structure modelling + programmatic assembly.

Follows the reconstruction-log pipeline, Stages D & E (which the original build
scripts performed together): map each role-tagged block to a docx primitive and
assemble the document so that ONE source page == ONE page-block, with the first
element of each page carrying a hard page break (page alignment, Section 5).
Running heads are inline paragraphs (not Word headers); footnotes are
bottom-of-page small text under a rule (not Word footnotes). The measured
typography (Stage G) is applied: face, size, body leading, and page margins.

Blocks understood (see 02_transcribe.py for the contract):
    title, heading, running_head, body (text|runs, bold_lead), list_item,
    table (rows, header, borders, shading), footnote, image (bbox, caption),
    form (lines)

Images are reproduced by cropping the high-resolution page render (recorded in
manifest.json) at the model's normalised bbox and inserting that crop in reading
order — this works for scanned and born-digital PDFs alike.

Uses python-docx (pure Python; replaces the original Node `docx` builder).

Usage:
    python3 05_assemble.py [--workdir work] [--out out/output.docx]

Input :  work/corrected.json, work/typography.json, work/manifest.json
Output:  out/output.docx
"""
import argparse, json, os, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

EMU_PER_PT = 12700

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "centre": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "justified": WD_ALIGN_PARAGRAPH.JUSTIFY}
HIGHLIGHT = {"yellow": WD_COLOR_INDEX.YELLOW, "green": WD_COLOR_INDEX.BRIGHT_GREEN,
             "cyan": WD_COLOR_INDEX.TURQUOISE, "pink": WD_COLOR_INDEX.PINK,
             "red": WD_COLOR_INDEX.RED, "gray": WD_COLOR_INDEX.GRAY_25, "grey": WD_COLOR_INDEX.GRAY_25}

# A scanned page that the model returns almost no text for (a cover, plate, photo
# page or full-page diagram) is reproduced faithfully by embedding the whole page
# render as one image rather than a few sparse text lines.
PLATE_TEXT_MAX = int(os.environ.get("PLATE_TEXT_MAX", "120"))


def load_json(path):
    """Read a JSON file, tolerating legacy cp1252/latin-1 files (older runs wrote
    non-UTF-8 bytes for chars like the em dash)."""
    p = Path(path)
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return json.loads(p.read_text(encoding=enc))
        except UnicodeDecodeError:
            continue
    return json.loads(p.read_text(encoding="utf-8", errors="replace"))


# ----------------------------------------------------------------------
# low-level helpers
# ----------------------------------------------------------------------
def set_default_font(doc, face, size_pt):
    style = doc.styles["Normal"]
    style.font.name = face
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), face)
    # zero the default inter-paragraph spacing — Word's 8-10pt "space after"
    # otherwise inflates every paragraph and overflows the page (parity drift).
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def set_page(section, w_pt, h_pt, margins_tw):
    section.page_width = Pt(w_pt)
    section.page_height = Pt(h_pt)
    clamp = lambda tw: min(max(tw, 432), 3600)  # 0.3in .. 2.5in
    section.top_margin = Pt(clamp(margins_tw["top"]) / 20)
    section.bottom_margin = Pt(clamp(margins_tw["bottom"]) / 20)
    section.left_margin = Pt(clamp(margins_tw["left"]) / 20)
    section.right_margin = Pt(clamp(margins_tw["right"]) / 20)


def body_leading(p, leading_pt):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(leading_pt)


def add_top_border(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "4"); top.set(qn("w:color"), "000000")
    pbdr.append(top); pPr.append(pbdr)


def set_table_borders(table, edges=("top", "left", "bottom", "right", "insideH", "insideV")):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)
    # OOXML requires tblBorders before shd/tblLayout/tblCellMar/tblLook in tblPr
    after = {qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"),
             qn("w:tblLook"), qn("w:tblCaption"), qn("w:tblDescription")}
    ref = next((c for c in tblPr if c.tag in after), None)
    if ref is not None:
        ref.addprevious(borders)
    else:
        tblPr.append(borders)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def set_cell_text_direction(cell, val="btLr"):
    """Rotate cell text (btLr = bottom-to-top) to mimic printed vertical headers."""
    tcPr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), val)
    tcPr.append(td)


def fix_settings(doc):
    """python-docx's default settings.xml has <w:zoom> without the required percent."""
    z = doc.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None:
        z.set(qn("w:percent"), "100")


def add_runs(p, blk):
    """Render a body block: prefer an explicit "runs" list (inline styling),
    else fall back to bold_lead + plain text. Returns nothing."""
    runs = blk.get("runs")
    if isinstance(runs, list) and runs:
        for rn in runs:
            r = p.add_run(rn.get("text", ""))
            if rn.get("b"): r.bold = True
            if rn.get("i"): r.italic = True
            if rn.get("u"): r.underline = True
            if rn.get("sup"): r.font.superscript = True
            if rn.get("sub"): r.font.subscript = True
            if rn.get("strike"): r.font.strike = True
            col = rn.get("color")
            if col:
                try:
                    r.font.color.rgb = RGBColor.from_string(str(col).lstrip("#"))
                except (ValueError, TypeError):
                    pass
            hl = rn.get("highlight")
            if hl and str(hl).lower() in HIGHLIGHT:
                r.font.highlight_color = HIGHLIGHT[str(hl).lower()]
        return
    lead = blk.get("bold_lead")
    if lead:
        rb = p.add_run(lead.rstrip() + " ")
        rb.bold = True
    p.add_run(blk.get("text", ""))


# ----------------------------------------------------------------------
# block -> primitive
# ----------------------------------------------------------------------
def emit_title(doc, blk, base):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(base * 0.4)
    r = p.add_run(blk.get("text", ""))
    r.bold = True
    r.font.size = Pt(base * (1.8 if blk.get("size") == "display" else 1.35))
    return p


def emit_heading(doc, blk, base):
    p = doc.add_paragraph()
    lvl = int(blk.get("level", 1))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if lvl <= 2 else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(base * 0.3); pf.space_after = Pt(base * 0.2)
    r = p.add_run(blk.get("text", ""))
    r.bold = True
    r.font.size = Pt(base * {1: 1.3, 2: 1.15}.get(lvl, 1.0))
    return p


def emit_running_head(doc, blk, base, page_index, text_width_pt):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    num = blk.get("page_number")
    title = blk.get("text", "")
    # book convention: even page -> number left, title centred; odd -> title centred, number right
    pf.tab_stops.add_tab_stop(Pt(text_width_pt / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Pt(text_width_pt), WD_TAB_ALIGNMENT.RIGHT)
    if page_index % 2 == 0:
        line = f"{num or ''}\t{title}"
    else:
        line = f"\t{title}\t{num or ''}"
    r = p.add_run(line)
    r.bold = True
    r.font.size = Pt(base * 0.9)
    return p


_ENUM = re.compile(r"\(([0-9]{1,3}|[a-zA-Z]{1,3})\)")


def _roman(s):
    vals = {"i": 1, "v": 5, "x": 10, "l": 50, "c": 100, "d": 500, "m": 1000}
    s = s.lower()
    if not s or any(ch not in vals for ch in s):
        return None
    total = prev = 0
    for ch in reversed(s):
        v = vals[ch]
        total += -v if v < prev else v
        prev = max(prev, v)
    return total


def _consecutive(nums):
    return len(nums) >= 2 and all(nums[i + 1] - nums[i] == 1 for i in range(len(nums) - 1))


def _seq_ok(tokens):
    """True if the labels form a sequential series under some reading
    (1,2,3 / i,ii,iii / a,b,c) — guards against splitting stray parentheticals."""
    cands = []
    if all(t.isdigit() for t in tokens):
        cands.append([int(t) for t in tokens])
    rom = [_roman(t) for t in tokens]
    if all(r is not None for r in rom):
        cands.append(rom)
    if all(len(t) == 1 and t.isalpha() for t in tokens):
        cands.append([ord(t) for t in tokens])
    return any(_consecutive(c) for c in cands)


def split_inline_list(text):
    """If a body block begins with a run of list markers ('(a) X (b) Y ...'),
    return (items, remainder): the maximal LEADING sequential run as
    [(label, text), ...] plus the trailing text after it (which the caller can
    recurse on). Returns None when the block does not start with a sequential
    marker run, so numbered clauses and inline references are left intact."""
    t = (text or "").strip()
    m = list(_ENUM.finditer(t))
    if len(m) < 2 or m[0].start() != 0:
        return None
    toks = [mk.group(1).lower() for mk in m]
    k = 0
    for k_try in range(len(toks), 1, -1):      # longest sequential prefix, >=2
        if _seq_ok(toks[:k_try]):
            k = k_try
            break
    if k < 2:
        return None
    items = []
    for i in range(k):
        s = m[i].end()
        e = m[i + 1].start() if i + 1 < k else (m[k].start() if k < len(m) else len(t))
        items.append((m[i].group(0), t[s:e].strip()))
    remainder = t[m[k].start():].strip() if k < len(m) else ""
    return items, remainder


def _list_paragraph(doc, label, text, base, leading, indent):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    left = base * (1.4 + 1.6 * indent)
    pf = p.paragraph_format
    pf.left_indent = Pt(left)
    pf.first_line_indent = Pt(-(base * 1.6))   # hanging
    pf.tab_stops.add_tab_stop(Pt(left), WD_TAB_ALIGNMENT.LEFT)
    p.add_run(f"{label}\t{text}")
    body_leading(p, leading)
    return p


def emit_body(doc, blk, base, leading):
    # rescue a list the model flattened/merged into one block: peel the leading
    # marker run into list items and recurse on the remainder.
    res = (None if (blk.get("runs") or blk.get("bold_lead"))
           else split_inline_list(blk.get("text", "")))
    if res:
        items, remainder = res
        first = None
        for label, itext in items:
            p = _list_paragraph(doc, label, itext, base, leading, 0)
            first = first or p
        if remainder.strip():
            p = emit_body(doc, {"type": "body", "text": remainder}, base, leading)
            first = first or p
        return first
    p = doc.add_paragraph()
    p.alignment = ALIGN.get(blk.get("align"), WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.first_line_indent = Pt(base * 1.4)
    add_runs(p, blk)
    body_leading(p, leading)
    return p


def emit_list_item(doc, blk, base, leading):
    return _list_paragraph(doc, blk.get("label", ""), blk.get("text", ""),
                           base, leading, int(blk.get("indent", 1)))


def _normalize_rows(rows):
    """Coerce whatever the model put in "rows" into a clean rectangular
    list-of-lists of strings. Vision models sometimes emit dict rows (e.g. a
    stray list_item) or ragged rows; this rescues them instead of crashing."""
    if not isinstance(rows, list):
        return []
    norm = []
    for r in rows:
        if isinstance(r, list):
            norm.append([str(c) for c in r])
        elif isinstance(r, dict):
            if "label" in r or "text" in r:
                norm.append([str(r.get("label", "")), str(r.get("text", ""))])
            else:
                norm.append([str(v) for v in r.values()])
        else:
            norm.append([str(r)])
    ncols = max((len(r) for r in norm), default=0)
    for r in norm:
        if len(r) < ncols:
            r.extend([""] * (ncols - len(r)))
    return norm


def _border_edges(blk):
    """Which table borders to draw, from blk['rules'] (grid|ledger|rows|box) or
    the legacy blk['borders'] bool. Ledger = column rules + outer box, no row
    rules (the classic register/ledger look)."""
    rules = str(blk.get("rules") or "").lower()
    if rules == "ledger":
        return ("top", "left", "bottom", "right", "insideV")
    if rules in ("rows", "horizontal"):
        return ("top", "left", "bottom", "right", "insideH")
    if rules == "box":
        return ("top", "left", "bottom", "right")
    if rules == "grid" or blk.get("borders"):
        return ("top", "left", "bottom", "right", "insideH", "insideV")
    return ()


def emit_table(doc, blk, base, text_width_pt):
    rows = _normalize_rows(blk.get("rows", []))
    if not rows:
        return None
    ncols = max((len(r) for r in rows), default=0)
    if ncols == 0:
        return None
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # column widths: explicit col_widths (fractions, e.g. from the refine pass)
    # win; otherwise proportional to the longest cell per column (8% floor).
    cw = blk.get("col_widths")
    if isinstance(cw, list) and len(cw) == ncols and sum(float(x) for x in cw) > 0:
        raw = [max(float(x), 0.02) for x in cw]
    else:
        colmax = [1] * ncols
        for row in rows:
            for ci in range(min(len(row), ncols)):
                colmax[ci] = max(colmax[ci], len(str(row[ci])))
        floor = 0.08 * sum(colmax)
        raw = [max(cm, floor) for cm in colmax]
    scale = text_width_pt / sum(raw)
    col_emu = [Emu(int(w * scale * EMU_PER_PT)) for w in raw]
    for ci, col in enumerate(table.columns):
        col.width = col_emu[ci]
    # border style: "rules" (grid|ledger|rows|box) wins; else borders=true -> grid
    edges = _border_edges(blk)
    if edges:
        set_table_borders(table, edges)
    header = blk.get("header", False)
    shade = blk.get("shading")          # "#DDD" (header) OR {row_index: "#DDD"}
    vertical = bool(blk.get("vertical_header"))  # headers printed rotated/vertical
    # text-fit levers (point C): the refine loop tunes these by RENDERING.
    ratio = float(blk.get("cell_size_ratio", 0.8))   # shrink font to fit
    cell_size = Pt(base * ratio)
    rh = blk.get("row_height")                        # taller rows for vertical headers
    row_h = Pt(float(rh)) if rh else Pt(base * ratio * 1.7)
    for r in table.rows:
        r.height = row_h
        r.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    def shade_for(ri):
        if isinstance(shade, dict):
            return shade.get(str(ri)) or shade.get(ri)
        if isinstance(shade, str):
            return shade if (header and ri == 0) else None
        return None

    for ri, row in enumerate(rows):
        fill = shade_for(ri)
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.width = col_emu[ci]
            txt = row[ci] if ci < len(row) else ""
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            run = cp.add_run(str(txt))
            run.font.size = cell_size
            if fill:
                set_cell_shading(cell, fill)
            if header and ri == 0:
                run.bold = True
                if vertical:
                    set_cell_text_direction(cell, "btLr")
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # merged cells: [{r,c,rowspan,colspan}] -> rectangular span (refine/model)
    for m in blk.get("merges") or []:
        try:
            r, c = int(m["r"]), int(m["c"])
            br = r + int(m.get("rowspan", 1)) - 1
            bc = c + int(m.get("colspan", 1)) - 1
            if (r, c) != (br, bc):
                table.cell(r, c).merge(table.cell(br, bc))
        except Exception:
            pass   # out-of-range / bad span: skip, never crash the build
    return table


def emit_footnote(doc, blk, base):
    p = doc.add_paragraph()
    add_top_border(p)
    marker = blk.get("marker", "")
    if marker:
        rs = p.add_run(marker)
        rs.font.superscript = True
        rs.font.size = Pt(base * 0.8)
        p.add_run("  ")
    r = p.add_run(blk.get("text", ""))
    r.font.size = Pt(base * 0.8)
    return p


def emit_form(doc, blk, leading):
    """A form/letter template: keep each printed line on its own line, left
    aligned, no justification or first-line indent (preserves field layout)."""
    lines = blk.get("lines") or ([blk.get("text", "")] if blk.get("text") else [])
    first = None
    for ln in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        body_leading(p, leading)
        p.add_run(ln)
        if first is None:
            first = p
    return first


def _norm_bbox(bbox):
    """Coerce a model bbox to four fractions in [0,1] as (x0,y0,x1,y1) or None."""
    if not (isinstance(bbox, (list, tuple)) and len(bbox) == 4):
        return None
    vals = []
    for v in bbox:
        try:
            v = float(v)
        except (TypeError, ValueError):
            return None
        if v > 1.0:            # tolerate 0..100 percentages
            v = v / 100.0
        vals.append(min(max(v, 0.0), 1.0))
    x0, y0, x1, y1 = vals
    if x1 <= x0 or y1 <= y0:
        return None
    return x0, y0, x1, y1


def emit_image(doc, blk, base, leading, page_w_pt, text_width_pt, page_img_path, media_dir, tag):
    """Crop the page render at the block's bbox and insert it as a picture.
    Falls back to a bracketed caption if the crop is not possible."""
    bbox = _norm_bbox(blk.get("bbox"))
    caption = blk.get("caption")
    saved = None
    if bbox and page_img_path and Path(page_img_path).is_file():
        try:
            im = Image.open(page_img_path)
            W, H = im.size
            x0, y0, x1, y1 = bbox
            crop = im.crop((int(x0 * W), int(y0 * H), int(x1 * W), int(y1 * H)))
            media_dir.mkdir(parents=True, exist_ok=True)
            saved = media_dir / f"{tag}.png"
            crop.save(saved)
            width_pt = min((x1 - x0) * page_w_pt, text_width_pt)
        except Exception as e:
            print(f"    ! image crop failed ({e}); inserting caption only")
            saved = None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if saved:
        p.add_run().add_picture(str(saved), width=Emu(int(width_pt * EMU_PER_PT)))
    else:
        body_leading(p, leading)
        r = p.add_run(f"[image: {caption}]" if caption else "[image]")
        r.italic = True
    if caption and saved:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption); cr.italic = True; cr.font.size = Pt(base * 0.85)
    return p


def page_text_len(page):
    """Total characters of real text the model captured on a page (used to spot
    image-dominant pages the model under-transcribed)."""
    n = 0
    for b in page["blocks"]:
        t = b.get("type")
        if t == "table":
            for r in _normalize_rows(b.get("rows", [])):
                n += sum(len(c) for c in r)
        elif t == "form":
            n += sum(len(x) for x in b.get("lines", []))
        elif t == "image":
            continue
        else:
            if b.get("runs"):
                n += sum(len(rn.get("text", "")) for rn in b["runs"])
            n += len(b.get("text", "")) + len(b.get("bold_lead") or "") + len(b.get("label") or "")
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir", default="work")
    ap.add_argument("--out", default="out/output.docx")
    args = ap.parse_args()
    work = Path(args.workdir)
    corrected = load_json(work / "corrected.json")
    typo = load_json(work / "typography.json")
    manifest = load_json(work / "manifest.json")

    base = float(typo["size_pt"])
    leading = float(typo["leading_pt"])
    face = typo["font_family"]
    margins = typo["margins_twips"]
    w_pt, h_pt = manifest["page_w_pt"], manifest["page_h_pt"]

    # page index -> high-res render path, text-layer flag and detected rotation
    page_img = {p["index"]: p.get("image") for p in manifest["pages"]}
    has_text = {p["index"]: p.get("has_text_layer", False) for p in manifest["pages"]}
    rotation = {p["index"]: int(p.get("rotation", 0) or 0) for p in manifest["pages"]}

    media_dir = Path(args.out).parent / "media"

    doc = Document()
    set_default_font(doc, face, base)

    skipped = 0
    plates = 0
    landscapes = 0
    # one Word section per source page (page alignment) so each page can carry
    # its own orientation; pages the pre-pass found rotated 90/270 go landscape.
    for pi, page in enumerate(corrected["pages"]):
        land = rotation.get(page["page"], 0) in (90, 270)
        pw, ph = (h_pt, w_pt) if land else (w_pt, h_pt)
        tw = pw - (margins["left"] + margins["right"]) / 20

        section = doc.sections[0] if pi == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        set_page(section, pw, ph, margins)
        section.orientation = WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        if land:
            landscapes += 1

        img_n = 0
        blocks = page["blocks"]
        # image-dominant scanned page -> embed the whole page render so covers/
        # plates/full-page figures are not lost. Either forced (page.plate, e.g.
        # from the refine pass) or auto-detected from sparse text.
        structural = any(b.get("type") in ("table", "form", "image") for b in blocks)
        auto_plate = (not has_text.get(page["page"], False) and not structural
                      and page_text_len(page) < PLATE_TEXT_MAX)
        if page.get("plate") or auto_plate:
            blocks = [{"type": "image", "bbox": [0, 0, 1, 1], "caption": None}]
            plates += 1
        for blk in blocks:
            t = blk.get("type")
            try:
                if t == "title":
                    emit_title(doc, blk, base)
                elif t == "heading":
                    emit_heading(doc, blk, base)
                elif t == "running_head":
                    emit_running_head(doc, blk, base, page["page"], tw)
                elif t == "body":
                    emit_body(doc, blk, base, leading)
                elif t == "list_item":
                    emit_list_item(doc, blk, base, leading)
                elif t == "table":
                    emit_table(doc, blk, base, tw)
                elif t == "footnote":
                    emit_footnote(doc, blk, base)
                elif t == "form":
                    emit_form(doc, blk, leading)
                elif t == "image":
                    img_n += 1
                    emit_image(doc, blk, base, leading, pw, tw,
                               page_img.get(page["page"]), media_dir,
                               f"fig-p{page['page']:04d}-{img_n}")
            except Exception as e:
                # one malformed block must never lose the whole document
                skipped += 1
                print(f"    ! page {page['page']} {t} block skipped: {e}")

    fix_settings(doc)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(f"Stage D+E done: {len(corrected['pages'])} page-blocks, "
          f"{face} {base}pt, leading {leading}pt -> {args.out}"
          + (f"  ({landscapes} landscape)" if landscapes else "")
          + (f"  ({plates} full-page image(s))" if plates else "")
          + (f"  ({skipped} block(s) skipped)" if skipped else ""))


if __name__ == "__main__":
    main()
